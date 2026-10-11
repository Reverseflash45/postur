"""Bangkitkan formulir cetak untuk pengambilan data (butuh python-docx).

    pip install python-docx
    python docs/buat_formulir.py

Keluaran di folder docs/: formulir-persetujuan.docx dan lembar-observasi.docx.
Bagian bertanda titik-titik diisi tangan atau disunting di Word sebelum dicetak.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

DOCS = Path(__file__).resolve().parent
JUDUL = ("Sistem Deteksi Kesalahan Postur Latihan Squat dan Push-Up Berbasis Pose "
         "Estimation dan Klasifikasi Deret Waktu pada Perangkat Android")
PENELITI = "Rafi Fernandito Setiawan"
INSTANSI = "Program Studi D4 Teknik Informatika, Fakultas Vokasi, Universitas Airlangga"
TITIK = "." * 48


def dokumen() -> Document:
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.5)
    gaya = d.styles["Normal"]
    gaya.font.name = "Times New Roman"
    gaya.font.size = Pt(10.5)
    gaya.paragraph_format.space_after = Pt(2)
    gaya.paragraph_format.line_spacing = 1.0
    return d


def judul(d: Document, teks: str, sub: str | None = None) -> None:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(teks)
    r.bold = True
    r.font.size = Pt(13)
    if sub:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(sub).italic = True


def bagian(d: Document, teks: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run(teks).bold = True


def lebar_kolom(t, lebar_cm: list[float]) -> None:
    """python-docx butuh lebar di setiap sel supaya Word menurutinya."""
    t.autofit = False
    for row in t.rows:
        for c, w in zip(row.cells, lebar_cm):
            c.width = Cm(w)


def isian(d: Document, label: list[str], lebar_label: float = 4.5) -> None:
    """Baris "Label : ......" yang titik duanya sejajar."""
    t = d.add_table(rows=len(label), cols=2)
    for i, a in enumerate(label):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = ": " + "." * 60
    lebar_kolom(t, [lebar_label, 16 - lebar_label])


def formulir_persetujuan() -> Path:
    d = dokumen()
    judul(d, "FORMULIR PERSETUJUAN SETELAH PENJELASAN", "(Informed Consent)")
    t = d.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (a, b) in enumerate([
        ("Judul penelitian", JUDUL),
        ("Peneliti", f"{PENELITI}\n{INSTANSI}"),
        ("Dosen pembimbing", TITIK),
        ("Kontak peneliti", TITIK),
    ]):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = b
    lebar_kolom(t, [4, 11])

    bagian(d, "Penjelasan untuk calon subjek")
    poin = [
        ("Tujuan", "Penelitian ini mengembangkan aplikasi Android yang menghitung repetisi squat dan push-up "
                   "melalui kamera ponsel dan menilai apakah setiap repetisi dilakukan dengan benar."),
        ("Prosedur", "Anda akan melakukan pemanasan ringan, lalu beberapa set squat dan/atau push-up tanpa beban, "
                     "masing-masing sekitar 10 repetisi. Pada sebagian set Anda diminta memperagakan kesalahan "
                     "ringan secara sengaja, misalnya turun kurang dalam. Di antara set tersedia waktu istirahat. "
                     "Seluruh kegiatan berlangsung sekitar 20–30 menit."),
        ("Data yang direkam", "Aplikasi hanya menyimpan koordinat titik tubuh (misalnya posisi lutut dan bahu) yang "
                              "dihitung langsung di ponsel. Video dan foto tidak disimpan. Anda dicatat dengan kode "
                              "anonim beserta rentang usia, jenis kelamin, pengalaman latihan, dan tinggi badan."),
        ("Risiko", "Risikonya setara dengan latihan ringan tanpa beban, misalnya rasa lelah atau pegal. Peragaan "
                   "kesalahan dibuat ringan dan tidak dipaksakan. Anda dapat berhenti kapan saja."),
        ("Manfaat", "Anda tidak memperoleh manfaat langsung. Hasil penelitian diharapkan membantu orang berlatih "
                    "dengan postur yang lebih aman tanpa pendamping."),
        ("Kerahasiaan", "Nama Anda hanya tercantum pada formulir ini, yang disimpan terpisah dari data penelitian. "
                        "Hasil penelitian dilaporkan dalam bentuk gabungan tanpa identitas perorangan."),
        ("Kesukarelaan", "Keikutsertaan bersifat sukarela. Anda berhak menolak atau mengundurkan diri kapan pun tanpa "
                         "konsekuensi apa pun, dan seluruh data Anda akan dihapus."),
    ]
    for a, b in poin:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(f"{a}. ").bold = True
        p.add_run(b)

    bagian(d, "Pernyataan subjek")
    d.add_paragraph("Saya yang bertanda tangan di bawah ini menyatakan bahwa:")
    for teks in [
        "saya telah membaca atau dibacakan penjelasan di atas, dan telah diberi kesempatan bertanya;",
        "saya berusia 18–45 tahun dan dalam kondisi sehat;",
        "saya tidak mengalami cedera lutut, pinggul, punggung, bahu, atau pergelangan tangan dalam enam bulan terakhir, "
        "tidak memiliki keluhan nyeri sendi, kondisi jantung, atau kondisi lain yang membatasi aktivitas fisik ringan, "
        "dan tidak sedang hamil;",
        "saya bersedia secara sukarela menjadi subjek penelitian ini.",
    ]:
        d.add_paragraph(f"☐  {teks}")

    isian(d, ["Nama", "Tanggal"], 3)
    t = d.add_table(rows=3, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.cell(0, 0).text = "Subjek,"
    t.cell(0, 1).text = "Peneliti,"
    t.cell(1, 0).text = "\n\n"
    t.cell(2, 0).text = "(" + "." * 32 + ")"
    t.cell(2, 1).text = f"({PENELITI})"
    for c in (t.cell(0, 0), t.cell(0, 1), t.cell(2, 0), t.cell(2, 1)):
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run("Diisi peneliti — kode subjek: ").bold = True
    p.add_run("." * 16 + "   (kode ini yang dimasukkan ke aplikasi, bukan nama)")

    keluar = DOCS / "formulir-persetujuan.docx"
    d.save(keluar)
    return keluar


KELAS = [
    ("SQ-0", "Squat · Benar"), ("SQ-1", "Squat · Kurang dalam"),
    ("SQ-2", "Squat · Lutut masuk ke dalam"), ("SQ-3", "Squat · Badan terlalu condong"),
    ("PU-0", "Push-up · Benar"), ("PU-1", "Push-up · Pinggul turun"),
    ("PU-2", "Push-up · Pinggul terangkat"), ("PU-3", "Push-up · Kurang dalam"),
]


def lembar_observasi() -> Path:
    d = dokumen()
    judul(d, "LEMBAR OBSERVASI PENGAMBILAN DATA")
    isian(d, ["Kode subjek", "Tanggal / jam", "Lokasi", "Pencahayaan", "Pakaian / alas kaki",
              "Ponsel & kamera", "Baris urutan Williams"])

    bagian(d, "Rekaman per set (isi sesuai urutan di aplikasi)")
    kolom = ["No.", "Kode", "Kelas", "Hitungan manual", "Hitungan otomatis", "Status", "Catatan"]
    t = d.add_table(rows=1 + len(KELAS), cols=len(kolom))
    t.style = "Table Grid"
    for j, k in enumerate(kolom):
        c = t.cell(0, j)
        c.text = k
        c.paragraphs[0].runs[0].bold = True
    for i, (kode, nama) in enumerate(KELAS, start=1):
        t.cell(i, 1).text = kode
        t.cell(i, 2).text = nama
        t.cell(i, 5).text = "✓ / ulang / lewat"
    lebar_kolom(t, [1.4, 1.3, 4.4, 1.9, 1.9, 2.5, 2.6])
    for row in t.rows:
        row.height = Cm(0.85)
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)

    bagian(d, "Keluhan, insiden, atau penghentian")
    for _ in range(4):
        d.add_paragraph("." * 110)

    bagian(d, "Pemeriksaan setelah sesi")
    for teks in [
        "Semua set sudah ditinjau per repetisi di aplikasi (Tinjau).",
        "Repetisi yang salah peraga sudah dibuang atau diberi label ulang.",
        "Formulir persetujuan sudah ditandatangani dan disimpan terpisah.",
    ]:
        d.add_paragraph(f"☐  {teks}")

    keluar = DOCS / "lembar-observasi.docx"
    d.save(keluar)
    return keluar


if __name__ == "__main__":
    for p in (formulir_persetujuan(), lembar_observasi()):
        print("ditulis:", p)
